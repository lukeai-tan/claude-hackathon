import importlib.util
import io
import tempfile
import unittest
from pathlib import Path
from PIL import Image
from docx import Document
from core import PDF


@unittest.skipUnless(importlib.util.find_spec('pdf2docx'), 'pdf2docx required; exercised by Windows CI')
class WordTests(unittest.TestCase):
    def test_editable_styles_images_and_table(self):
        with tempfile.TemporaryDirectory() as folder:
            pdf = PDF()
            p = pdf.doc.new_page()
            p.insert_text((50, 60), 'Editable heading', fontsize=22, fontname='hebo', color=(0.2, 0.3, 0.7))
            p.insert_text((50, 90), 'Body text remains editable.', fontsize=12)
            for x in (50, 200, 350):
                p.draw_line((x, 150), (x, 230), color=(0, 0, 0))
            for y in (150, 190, 230):
                p.draw_line((50, y), (350, y), color=(0, 0, 0))
            for x, y, text in [(60, 175, 'Item'), (210, 175, 'Count'), (60, 215, 'Apples'), (210, 215, '3')]:
                p.insert_text((x, y), text, fontsize=12)
            image = io.BytesIO()
            Image.new('RGB', (80, 60), 'red').save(image, format='PNG')
            p.insert_image((50, 280, 130, 340), stream=image.getvalue())
            output = Path(folder) / 'layout.docx'
            try:
                pdf.export_docx(output)
            finally:
                pdf.close()
            doc = Document(output)
            runs = [r for paragraph in doc.paragraphs for r in paragraph.runs]
            self.assertTrue(any('Editable heading' in r.text for r in runs))
            self.assertTrue(any(r.font.size and r.font.size.pt >= 20 for r in runs))
            self.assertGreaterEqual(len(doc.tables), 1)
            self.assertIn('Apples', '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
            self.assertGreaterEqual(len(doc.inline_shapes), 1)

    def test_scan_requires_ocr(self):
        pdf = PDF()
        pdf.doc.new_page()
        try:
            with self.assertRaisesRegex(ValueError, 'OCR'):
                pdf.export_docx('should-not-exist.docx')
        finally:
            pdf.close()
