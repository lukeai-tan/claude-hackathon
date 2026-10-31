"""Exercise real Tk selection and replacement controls in the Windows EXE."""
import tempfile
from pathlib import Path
import fitz
import tkinter as tk
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
from text_tools import find_matches


def exercise_app(app):
    app.new()
    p = app.pdf.doc[0]
    p.insert_text((50, 100), 'problem one and problem two', fontsize=18, fontname='hebo')
    app.refresh()
    app.update()
    r = fitz.Rect(find_matches(app.pdf.doc, 'problem')[1]['rect'])
    app.canvas.event_generate('<ButtonPress-1>', x=int(r.x0 + 21), y=int(r.y0 + 22))
    app.canvas.event_generate('<B1-Motion>', x=int(r.x1 + 19), y=int(r.y0 + 22))
    app.update()
    assert app.text_quads, 'Mouse drag did not select text'
    app.highlight()
    assert len(list(app.pdf.doc[0].annots())) == 1
    app.replace()
    app.update()
    def descendants(widget):
        for child in widget.winfo_children():
            yield child
            yield from descendants(child)
    controls = list(descendants(app))
    entries = [w for w in controls if w.winfo_class() == 'TEntry']
    entries[0].insert(0, 'problem')
    entries[1].insert(0, 'fixed')
    next(w for w in controls if w.winfo_class() == 'TButton' and w.cget('text') == 'Find matches').invoke()
    app.update()
    # Select second match through the list and invoke actual replacement button.
    match_list = next(w for w in controls if w.winfo_class() == 'Listbox' and w is not app.pages)
    match_list.selection_clear(0, 'end')
    match_list.selection_set(1)
    match_list.event_generate('<<ListboxSelect>>')
    app.update()
    next(w for w in controls if w.winfo_class() == 'TButton' and w.cget('text') == 'Replace selected').invoke()
    app.update()
    assert len(app.pdf.doc[0].search_for('problem')) == 1
    assert len(app.pdf.doc[0].search_for('fixed')) == 1
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / 'packaged-export.docx'
        app.pdf.export_docx(path)
        doc = Document(path)
        editable_text = ''.join(node.text or '' for node in doc.element.xpath('.//w:t'))
        print('Packaged Word export editable text:', repr(editable_text))
        assert 'fixed' in editable_text, 'Replacement missing from editable Word text'

    # Close search/replace dialogs and exercise direct section editing.
    for child in app.winfo_children():
        if isinstance(child, tk.Toplevel):
            child.destroy()
    app.dirty = False
    app.new()
    p = app.pdf.doc[0]
    p.insert_text((50, 100), 'Editable text', fontname='hebo', fontsize=18)
    p.insert_text((50, 250), 'Keep me', fontsize=12)
    app.refresh()
    app.update()
    assert app.pages.images, 'Sidebar thumbnail was not rendered'
    app.edit_mode()
    assert app.canvas.find_withtag('text-box'), 'Edit mode has no text box outlines'
    app.canvas.event_generate('<ButtonPress-1>', x=80, y=115)
    app.update()
    assert app.inline is not None, 'Click did not open inline editor'
    editor = app.inline
    editor.widget.focus_force()
    editor.widget.mark_set('insert', '1.4')
    editor.widget.insert('insert', 'X')
    app.update()
    assert editor.widget.get('1.0', 'end-1c') == 'EditXable text'
    editor.widget.event_generate('<BackSpace>')
    app.update()
    assert editor.widget.get('1.0', 'end-1c') == 'Editable text', 'Backspace did not edit live'
    editor.widget.delete('1.0', 'end')
    editor.widget.insert('1.0', 'Live edit')
    editor.key_commit()
    app.update()
    assert 'Live edit' in app.pdf.doc[0].get_text()
    app.undo()
    assert 'Editable text' in app.pdf.doc[0].get_text()
    app.redo()
    assert 'Live edit' in app.pdf.doc[0].get_text()

    # Actual Ctrl+F opens find; results are page overlays rather than PDF edits.
    app.focus_force()
    app.event_generate('<Control-f>')
    app.update()
    assert app.search_window is not None, 'Ctrl+F did not open find'
    app.find_entry.insert(0, 'Live')
    app.after(300, app.quit)
    app.mainloop()
    app.update()
    assert app.canvas.find_withtag('search-result'), 'Find did not highlight words on page'
    zoom = app.zoom
    app.wheel(SimpleNamespace(state=4, delta=120, num=0, widget=app.canvas))
    assert app.zoom > zoom, 'Ctrl+wheel did not zoom'
    app.scale(1 / 1.1)

    # Drag selection offers contextual actions, including text redaction.
    app.selection_mode.set('text')
    app.clear_selection()
    app.render()
    app.canvas.xview_moveto(0)
    app.canvas.yview_moveto(0)
    r = app.pdf.doc[0].search_for('Live')[0]
    app.canvas.event_generate('<ButtonPress-1>', x=int(r.x0 + 21), y=int(r.y0 + 22))
    app.canvas.event_generate('<B1-Motion>', x=int(r.x1 + 19), y=int(r.y0 + 22))
    app.canvas.event_generate('<ButtonRelease-1>', x=int(r.x1 + 19), y=int(r.y0 + 22))
    app.update()
    assert app.context_frame is not None, 'Selection did not show contextual actions'
    labels = [w.cget('text') for w in app.context_frame.winfo_children()]
    assert all(s in labels for s in ('Edit section', 'Highlight', 'Redact', 'Copy'))
    from tkinter import messagebox
    with patch.object(messagebox, 'askyesno', return_value=True):
        app.redact()
    assert not app.pdf.doc[0].search_for('Live'), 'Selected text was not redacted'
    assert app.pdf.doc[0].search_for('Keep me'), 'Text redaction affected another section'
    app.undo()
    assert app.pdf.doc[0].search_for('Live')

    # Native PDF sections move by their border, and shorter edits compact.
    app.edit_mode()
    r = app.pdf.doc[0].search_for('Live edit')[0]
    app.canvas.event_generate('<ButtonPress-1>', x=int(r.x0 + 24), y=int(r.y0 + 25))
    app.update()
    editor = app.inline
    assert editor is not None
    original = fitz.Rect(editor.rect)
    x, y = int(original.x0 + 45), int(original.y0 + 16)
    app.canvas.event_generate('<Motion>', x=x, y=y)
    app.canvas.event_generate('<ButtonPress-1>', x=x, y=y)
    app.canvas.event_generate('<B1-Motion>', x=x + 60, y=y + 40)
    app.canvas.event_generate('<ButtonRelease-1>', x=x + 60, y=y + 40)
    app.update()
    assert editor.rect.x0 > original.x0, 'Native PDF box border was not draggable'
    editor.widget.delete('1.0', 'end')
    editor.widget.insert('1.0', '\nFew\n\n')
    editor.key_commit()
    app.update()
    moved = app.pdf.doc[0].search_for('Few')
    assert len(moved) == 1 and moved[0].x0 > original.x0
    assert not app.pdf.doc[0].search_for('Live edit'), 'Moving native text left a duplicate'
    app.undo()
    assert app.pdf.doc[0].search_for('Live edit')
    app.redo()
    assert app.pdf.doc[0].search_for('Few')

    # Create a box using an actual canvas drag, then format only a selection.
    from rich_boxes import boxes
    from tkinter import colorchooser
    for child in app.winfo_children():
        if isinstance(child, tk.Toplevel):
            child.destroy()
    app.search_window = None
    app.dirty = False
    app.new()
    app.canvas.xview_moveto(0)
    app.canvas.yview_moveto(0)
    app.text()
    app.canvas.event_generate('<ButtonPress-1>', x=80, y=320)
    app.canvas.event_generate('<B1-Motion>', x=380, y=450)
    app.canvas.event_generate('<ButtonRelease-1>', x=380, y=450)
    app.update()
    editor = app.inline
    assert editor is not None, 'Drag did not create a text box'
    editor.widget.insert('1.0', 'Red blue text')
    editor.widget.tag_add('sel', '1.0', '1.3')
    with patch.object(colorchooser, 'askcolor', return_value=((255, 0, 0), '#ff0000')):
        editor.choose_text()
    assert editor.style_at('1.1')['color'] == '#ff0000'
    assert editor.style_at('1.5')['color'] == '#000000', 'Colour escaped the selected text'
    editor.toggle('bold')
    editor.local_undo()
    assert not editor.style_at('1.1')['bold'], 'Local undo did not restore formatting'
    editor.local_redo()
    assert editor.style_at('1.1')['bold']
    with patch.object(colorchooser, 'askcolor', return_value=((255, 238, 221), '#ffeedd')):
        next(w for w in descendants(editor.toolbar) if w.winfo_class() == 'TButton' and w.cget('text') == 'Fill').invoke()
    assert editor.widget.cget('background') == '#ffeedd', 'Fill did not update live'
    editor.key_commit()
    app.update()
    assert len(boxes(app.pdf.doc[0])) == 1
    assert app.pdf.doc[0].search_for('Red blue text'), 'Box text is not searchable PDF text'

    # Click the saved box again, drag its top border, and resize its corner.
    app.canvas.event_generate('<ButtonPress-1>', x=100, y=340)
    app.update()
    editor = app.inline
    assert editor is not None, 'Saved box could not be reopened'
    original = fitz.Rect(editor.rect)
    x, y = int(original.x0 + 70), int(original.y0 + 18)
    app.canvas.event_generate('<Motion>', x=x, y=y)
    app.canvas.event_generate('<ButtonPress-1>', x=x, y=y)
    app.canvas.event_generate('<B1-Motion>', x=x + 30, y=y + 20)
    app.canvas.event_generate('<ButtonRelease-1>', x=x + 30, y=y + 20)
    app.update()
    assert editor.rect.x0 > original.x0, 'Dragging border did not move box'
    old_width = editor.rect.width
    x, y = round(editor.rect.x1 + 20), round(editor.rect.y1 + 20)
    app.canvas.event_generate('<Motion>', x=x, y=y)
    app.canvas.event_generate('<ButtonPress-1>', x=x, y=y)
    app.canvas.event_generate('<B1-Motion>', x=x + 30, y=y + 20)
    app.canvas.event_generate('<ButtonRelease-1>', x=x + 30, y=y + 20)
    app.update()
    assert editor.rect.width > old_width, 'Corner drag did not resize box'
    editor.key_commit()
    app.update()
    moved = boxes(app.pdf.doc[0])[0]
    app.undo()
    assert boxes(app.pdf.doc[0])[0]['rect'] == list(original), 'Undo did not restore box position'
    app.redo()
    assert boxes(app.pdf.doc[0])[0]['rect'] == moved['rect']

    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / 'styled-box.pdf'
        app.pdf.save(path)
        app.load(path)
        app.update()
        saved = boxes(app.pdf.doc[0])[0]
        assert saved['fill'] == '#ffeedd'
        assert saved['runs'][0]['style']['color'] == '#ff0000'
        # The same normal PDF text must reach editable Word output.
        word = Path(folder) / 'styled-box.docx'
        app.pdf.export_docx(word)
        xml_text = ''.join(n.text or '' for n in Document(word).element.xpath('.//w:t'))
        assert 'Red' in xml_text and 'blue text' in xml_text, 'Added box missing in Word'
        # Release the Windows file handle before TemporaryDirectory cleanup.
        app.pdf.restore(app.pdf.bytes())
        app.path = None

    # Reorder through real thumbnail drag/drop, and exercise undo/redo.
    app.pdf.doc.new_page().insert_text((50, 50), 'Page two')
    app.pdf.doc.new_page().insert_text((50, 50), 'Page three')
    app.refresh()
    app.update()
    tree = app.pages.tree
    first, third = tree.bbox('0'), tree.bbox('2')
    assert first and third, 'Thumbnail rows were not visible'
    start, end = first[1] + first[3] // 2, third[1] + third[3] - 8
    tree.event_generate('<ButtonPress-1>', x=70, y=start)
    tree.event_generate('<B1-Motion>', x=70, y=end)
    tree.event_generate('<ButtonRelease-1>', x=70, y=end)
    app.update()
    assert 'Page two' in app.pdf.doc[0].get_text(), 'Thumbnail drag did not reorder pages'
    assert boxes(app.pdf.doc[2]), 'Reordering lost the text box'
    assert app.index == 2, 'Moved page did not remain selected'
    app.undo()
    assert boxes(app.pdf.doc[0]), 'Undo did not restore page order'
    app.redo()
    assert boxes(app.pdf.doc[2]), 'Redo did not restore page order'
    print('Styled boxes, selection formatting, live fill, drag/resize, Word export and thumbnail reorder passed.')

    # Pen strokes preview live, retain colour/width, survive export, and undo.
    app.canvas.xview_moveto(0)
    app.canvas.yview_moveto(0)
    with patch.object(colorchooser, 'askcolor', return_value=((0, 128, 0), '#008000')):
        app.pen_color()
    app.ink_width.set('4')
    app.pen_mode()
    app.canvas.event_generate('<ButtonPress-1>', x=80, y=80)
    for x, y in [(100, 100), (120, 80), (140, 100), (160, 80)]:
        app.canvas.event_generate('<B1-Motion>', x=x, y=y)
    assert app.canvas.find_withtag('ink-preview'), 'Pen had no live preview'
    app.canvas.event_generate('<ButtonRelease-1>', x=170, y=90)
    app.update()
    assert not app.canvas.find_withtag('ink-preview')
    page = app.pdf.doc[app.index]
    ink = list(page.annots())
    assert len(ink) == 1 and ink[0].type[0] == fitz.PDF_ANNOT_INK
    assert ink[0].border['width'] == 4
    app.undo()
    assert not list(app.pdf.doc[app.index].annots() or [])
    app.redo()
    assert len(list(app.pdf.doc[app.index].annots())) == 1
    with tempfile.TemporaryDirectory() as folder:
        path = Path(folder) / 'ink.pdf'
        app.pdf.save(path)
        with fitz.open(path) as saved:
            page = saved[app.index]
            ink = list(page.annots())
            assert len(ink[0].vertices[0]) >= 5
            assert abs(ink[0].colors['stroke'][1] - 128 / 255) < .01
    app.canvas.event_generate('<ButtonPress-1>', x=200, y=100)
    app.escape_mode()
    app.canvas.event_generate('<ButtonRelease-1>', x=220, y=100)
    assert len(list(app.pdf.doc[app.index].annots())) == 1, 'Escape committed an unfinished stroke'
    app.escape_mode()
    assert app.selection_mode.get() == 'text'
    print('Freehand pen, live preview, export, colour/width, cancellation and undo/redo passed.')
    from gui_performance import exercise_performance
    exercise_performance(app)
